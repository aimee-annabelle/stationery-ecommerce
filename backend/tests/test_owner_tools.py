"""Stage 6: reports, activity log, staff & roles with invites, shop settings, owner alerts, daily summary."""

import io
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import ActivityLog, ActivityReview, Order, Setting, Staff, StaffInvite
from app.services import daily_summary, shop_settings
from tests.conftest import STAFF_PASSWORD, bearer, staff_token
from tests.test_catalog import last_log, product_id
from tests.test_print_requests import send, staff_id
from tests.test_shop import place
from tests.test_staff_orders import ORDERS, advance, online_order, pick_all, staff_id_of

OWNER_ONLY = [("GET", "/api/v1/staff/reports"), ("GET", "/api/v1/staff/activity"), ("GET", "/api/v1/staff/team"),
              ("GET", "/api/v1/staff/settings")]


@pytest.mark.parametrize("method, path", OWNER_ONLY)
def test_owner_tools_are_owner_only(client, seeded, headers, method, path):
    assert client.request(method, path, headers=headers["owner"]).status_code == 200
    assert client.request(method, path, headers=headers["manager"]).status_code == 403
    assert client.request(method, path, headers=headers["worker"]).status_code == 403


# Reports


def collect_print_job(client, db, headers, price=4500):
    rid = staff_id(client, send(client).json()["number"], headers["owner"])
    day = (shop_settings.today_in_kigali() + timedelta(days=3)).isoformat()
    client.post(f"/api/v1/staff/print-requests/{rid}/quote", json={"quote_rwf": price, "ready_on": day}, headers=headers["owner"])
    for body in ({}, {}, {"paid_with": "cash"}):
        client.post(f"/api/v1/staff/print-requests/{rid}/status", json=body, headers=headers["worker"])


def test_report_numbers(client, seeded, headers, team):
    online_order(client, seeded, "pay_at_pickup")  # 13,000 (2 notebooks + charger)
    cancelled = online_order(client, seeded, "pay_at_pickup", items=[{"product_id": product_id(seeded, "SCH-002"), "qty": 1}])
    client.post(f"{ORDERS}/{staff_id_of(client, seeded, cancelled['number'], headers['owner'])}/cancel", json={"reason": "x"}, headers=headers["owner"])
    client.post(ORDERS, json={"items": [{"product_id": product_id(seeded, "ELE-002"), "qty": 1}], "paid_with": "cash"}, headers=headers["worker"])  # 7,000
    collect_print_job(client, seeded, headers)  # 4,500

    report = client.get("/api/v1/staff/reports", params={"range": "7d"}, headers=headers["owner"]).json()

    assert report["revenue_rwf"] == 13000 + 7000 + 4500
    assert (report["orders"], report["average_order_rwf"], report["print_jobs_completed"]) == (2, 10000, 1)
    assert len(report["buckets"]) == 7 and report["buckets"][-1]["revenue_rwf"] == 24500
    cats = {c["name"]: c["revenue_rwf"] for c in report["categories"]}
    assert cats == {"School supplies": 5000, "Phone accessories": 8000, "Electronics": 7000, "Print & binding": 4500}
    assert report["top_products"][0] == {"name": "USB-C fast charger", "units": 1, "revenue_rwf": 8000}
    worker = next(s for s in report["staff"] if s["name"] == team["worker"].name)
    assert (worker["sales_rwf"], worker["print_jobs"], worker["orders_handled"]) == (7000, 1, 1)


def test_report_ranges(client, seeded, headers):
    month = client.get("/api/v1/staff/reports", params={"range": "30d"}, headers=headers["owner"]).json()
    year = client.get("/api/v1/staff/reports", params={"range": "year"}, headers=headers["owner"]).json()

    assert sum((datetime.fromisoformat(b["end"]) - datetime.fromisoformat(b["start"])).days + 1 for b in month["buckets"]) == 30
    assert len(year["buckets"]) == shop_settings.today_in_kigali().month


def test_report_exports(client, seeded, headers):
    online_order(client, seeded, "pay_at_pickup")

    xlsx = client.get("/api/v1/staff/reports/export", params={"format": "xlsx"}, headers=headers["owner"])
    pdf = client.get("/api/v1/staff/reports/export", params={"format": "pdf", "range": "year"}, headers=headers["owner"])

    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames == ["Summary", "Revenue", "Categories", "Top products", "Staff"]
    assert book["Summary"]["B2"].value == 13000
    assert pdf.content.startswith(b"%PDF") and pdf.headers["content-disposition"].endswith('.pdf"')


# Activity log


def test_activity_filters_wording_and_export(client, seeded, headers, team):
    pid = product_id(seeded, "OFF-003")
    client.patch(f"/api/v1/staff/products/{pid}", json={"price_rwf": 9800, "location": "Shelf B3"}, headers=headers["manager"])

    everything = client.get("/api/v1/staff/activity", headers=headers["owner"]).json()
    sensitive = client.get("/api/v1/staff/activity", params={"sensitive": True, "kind": "products"}, headers=headers["owner"]).json()
    by_worker = client.get("/api/v1/staff/activity", params={"staff_id": str(team["worker"].id), "kind": "products"}, headers=headers["owner"]).json()
    csv = client.get("/api/v1/staff/activity/export", params={"sensitive": True}, headers=headers["owner"]).text

    assert [i["what"] for i in sensitive["items"]] == ["changed the price of SKU OFF-003 from 9,500 RWF to 9,800 RWF"]
    assert sensitive["items"][0]["before"] == {"price_rwf": 9500} and sensitive["items"][0]["device"] == "pytest-browser"
    assert by_worker["items"] == []
    assert everything["total"] >= 4  # three sign-ins + the edits
    assert "changed the price of SKU OFF-003" in csv and csv.startswith("﻿When (Kigali)")


def test_looks_fine_records_a_review_without_changing_the_entry(client, seeded, headers):
    client.patch(f"/api/v1/staff/products/{product_id(seeded, 'OFF-003')}", json={"price_rwf": 9800}, headers=headers["manager"])
    eye = client.get("/api/v1/staff/dashboard", headers=headers["owner"]).json()["pocket"]["needs_eye"]
    entry = next(e for e in eye if e["action"] == "product.price_changed")

    assert client.post(f"/api/v1/staff/activity/{entry['id']}/reviewed", headers=headers["owner"]).status_code == 204
    client.post(f"/api/v1/staff/activity/{entry['id']}/reviewed", headers=headers["owner"])  # twice is fine

    after = client.get("/api/v1/staff/dashboard", headers=headers["owner"]).json()["pocket"]["needs_eye"]
    assert entry["id"] not in [e["id"] for e in after]
    assert seeded.scalar(select(ActivityLog.after).where(ActivityLog.id == entry["id"])) == {"price_rwf": 9800}
    assert len(seeded.scalars(select(ActivityReview)).all()) == 1


# Staff & roles


def invite(client, headers, **overrides):
    body = {"name": "Diane K.", "phone": "0788 000 004", "email": "diane@example.com", "role": "worker", **overrides}
    return client.post("/api/v1/staff/team/invite", json=body, headers=headers["owner"])


def test_invite_accept_and_sign_in(client, seeded, headers, outbox_mail):
    response = invite(client, headers)

    assert response.status_code == 201, response.text
    link = response.json()["link"]
    assert link.startswith("http://localhost:5173/staff/invite?token=")
    assert outbox_mail.sent[-1].to == "diane@example.com" and link in outbox_mail.sent[-1].body
    assert last_log(seeded, "staff.added").is_sensitive
    token = link.split("token=")[1]
    assert client.get("/api/v1/staff/auth/invite", params={"token": token}).json()["name"] == "Diane K."

    joined = client.post("/api/v1/staff/auth/accept-invite", json={"token": token, "password": "new-staff-pass"})
    again = client.post("/api/v1/staff/auth/accept-invite", json={"token": token, "password": "another-pass"})

    assert joined.status_code == 200 and joined.json()["staff"]["role"] == "worker"
    assert again.status_code == 404
    assert client.post("/api/v1/staff/auth/login", json={"email": "diane@example.com", "password": "new-staff-pass"}).status_code == 200


def test_expired_or_replaced_invites_fail(client, seeded, headers, outbox_mail):
    first = invite(client, headers).json()
    resent = client.post(f"/api/v1/staff/team/{first['member']['id']}/invite", headers=headers["owner"]).json()
    old_token, new_token = first["link"].split("token=")[1], resent["link"].split("token=")[1]

    assert client.post("/api/v1/staff/auth/accept-invite", json={"token": old_token, "password": "pass-pass-1"}).status_code == 404
    seeded.execute(StaffInvite.__table__.update().values(expires_at=datetime.now() - timedelta(days=1)))
    assert client.post("/api/v1/staff/auth/accept-invite", json={"token": new_token, "password": "pass-pass-1"}).status_code == 404


def test_duplicate_email_and_owner_rules(client, seeded, headers, team):
    assert invite(client, headers).status_code == 201
    assert invite(client, headers).json()["code"] == "email_taken"
    owner_id = str(team["owner"].id)
    assert client.patch(f"/api/v1/staff/team/{owner_id}", json={"role": "worker"}, headers=headers["owner"]).json()["code"] == "owner_protected"

    changed = client.patch(f"/api/v1/staff/team/{team['worker'].id}", json={"role": "manager"}, headers=headers["owner"]).json()
    removed = client.patch(f"/api/v1/staff/team/{team['worker'].id}", json={"is_active": False}, headers=headers["owner"]).json()
    assert changed["role"] == "manager" and removed["is_active"] is False
    listing = client.get("/api/v1/staff/team", headers=headers["owner"]).json()
    assert listing["members"][0]["role"] == "owner"
    assert "costs.view" not in listing["permissions"]["worker"] and "reports.view" in listing["permissions"]["owner"]


# Settings


def test_settings_change_is_checked_and_logged(client, seeded, headers):
    current = client.get("/api/v1/staff/settings", headers=headers["owner"]).json()["settings"]
    orders = {**current["orders"], "pickup_hold_days": 5}

    saved = client.patch("/api/v1/staff/settings", json={"orders": orders}, headers=headers["owner"])
    bad = client.patch("/api/v1/staff/settings", json={"orders": {**orders, "pickup_hold_days": 90}}, headers=headers["owner"])
    unknown = client.patch("/api/v1/staff/settings", json={"orders": {**orders, "surprise": 1}}, headers=headers["owner"])
    no_way_to_pay = client.patch("/api/v1/staff/settings", headers=headers["owner"], json={"payment_methods": {
        "momo": {"enabled": False, "merchant_code": ""}, "airtel": {"enabled": False, "merchant_code": ""}, "pay_at_pickup": {"enabled": False}}})

    assert saved.status_code == 200 and saved.json()["settings"]["orders"]["pickup_hold_days"] == 5
    entry = last_log(seeded, "settings.updated")
    assert entry.is_sensitive and entry.before == {"orders": {"pickup_hold_days": 3}} and entry.after == {"orders": {"pickup_hold_days": 5}}
    assert bad.json()["code"] == unknown.json()["code"] == no_way_to_pay.json()["code"] == "invalid_setting"
    assert "at least one way to pay" in no_way_to_pay.json()["detail"]


def test_hours_and_closures_drive_pickup_days(client, seeded, headers):
    tomorrow = (shop_settings.today_in_kigali() + timedelta(days=1)).isoformat()
    client.patch("/api/v1/staff/settings", json={"closures": [tomorrow]}, headers=headers["owner"])

    assert tomorrow not in client.get("/api/v1/settings/public").json()["pickup_days"]


# Owner alerts


def test_price_change_by_manager_emails_the_owner_after_commit(client, seeded, headers, team, outbox_mail):
    client.patch(f"/api/v1/staff/products/{product_id(seeded, 'OFF-003')}", json={"price_rwf": 9800}, headers=headers["manager"])

    alert = outbox_mail.sent[-1]
    assert alert.to == team["owner"].email and "Price changed" in alert.subject
    assert "Old price: 9,500 RWF. New price: 9,800 RWF." in alert.body


def test_no_alert_for_the_owners_own_change_or_when_switched_off(client, seeded, headers, outbox_mail):
    client.patch(f"/api/v1/staff/products/{product_id(seeded, 'OFF-003')}", json={"price_rwf": 9800}, headers=headers["owner"])
    assert outbox_mail.sent == []

    alerts = seeded.scalar(select(Setting).where(Setting.key == "owner_alerts"))
    alerts.value = {**alerts.value, "price": False}
    seeded.flush()
    client.patch(f"/api/v1/staff/products/{product_id(seeded, 'OFF-003')}", json={"price_rwf": 9900}, headers=headers["manager"])
    assert outbox_mail.sent == []


def test_failed_change_sends_no_alert(client, seeded, headers, outbox_mail):
    # Adjusting below zero fails after the movement is attempted; nothing is committed, so nothing is sent.
    response = client.post("/api/v1/staff/stock/adjust", headers=headers["manager"], json={
        "product_id": product_id(seeded, "OFF-003"), "change_qty": -50, "reason": "adjusted", "note": "count"})

    assert response.status_code == 400 and outbox_mail.sent == []


def test_alerts_go_to_the_configured_address(client, seeded, headers, outbox_mail):
    current = client.get("/api/v1/staff/settings", headers=headers["owner"]).json()["settings"]["owner_alerts"]
    client.patch("/api/v1/staff/settings", json={"owner_alerts": {**current, "email": "boss@example.com"}}, headers=headers["owner"])

    client.post("/api/v1/staff/stock/adjust", headers=headers["manager"], json={
        "product_id": product_id(seeded, "SCH-001"), "change_qty": -1, "reason": "damaged", "note": "torn"})

    assert outbox_mail.sent[-1].to == "boss@example.com" and "Stock adjusted" in outbox_mail.sent[-1].subject


# Daily summary and receipts


def test_daily_summary_sends_once_after_the_time(seeded, team, outbox_mail):
    today = shop_settings.today_in_kigali()
    early = datetime.combine(today, datetime.min.time(), tzinfo=shop_settings.KIGALI).replace(hour=23, minute=58)

    assert daily_summary.send_if_due(seeded, early) is False
    assert daily_summary.send_if_due(seeded, early.replace(minute=59)) is True
    assert daily_summary.send_if_due(seeded, early.replace(minute=59, second=30)) is False
    seeded.commit()

    assert len(outbox_mail.sent) == 1
    summary = outbox_mail.sent[0]
    assert summary.to == team["owner"].email and "daily summary" in summary.subject
    assert "Low stock (5):" in summary.body and "Heavy-duty stapler: 2 left" in summary.body


def test_receipt_email_after_payment_when_switched_on(client, seeded, headers, outbox_mail):
    receipt = seeded.scalar(select(Setting).where(Setting.key == "receipt"))
    receipt.value = {**receipt.value, "email_after_payment": True}
    seeded.flush()
    order = online_order(client, seeded)
    client.post(f"/api/v1/orders/{order['number']}/payment", json={"phone": "0788123456", "method": "momo", "transaction_id": "PAY12345"})
    oid = staff_id_of(client, seeded, order["number"], headers["worker"])

    client.post(f"{ORDERS}/{oid}/verify-payment", json={"found": True}, headers=headers["worker"])

    mail = [e for e in outbox_mail.sent if e.to == "aline@example.com" and "receipt" in e.subject.lower()]
    assert len(mail) == 1 and "13,000 RWF" in mail[0].body
