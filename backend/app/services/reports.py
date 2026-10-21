"""S5 Reports (owner only), with Excel and PDF exports.

Revenue = value of orders that weren't cancelled (online and walk-in, the same as "Sales today" on the
dashboard) + collected print jobs (paid at pickup). Days are counted in Kigali time.
"""

import io
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from sqlalchemy import func, literal_column, select
from sqlalchemy.orm import Session

from app.models import Category, Order, OrderItem, OrderStatusHistory, PrintRequest, PrintRequestHistory, Product, Staff
from app.services import shop_settings
from app.services.staff_orders import kigali_day_start

RANGES = {"7d": "last 7 days", "30d": "last 30 days", "year": "this year"}


@dataclass
class Report:
    range: str
    label: str
    start: date
    end: date
    revenue_rwf: int = 0
    orders: int = 0
    average_order_rwf: int = 0
    print_jobs_completed: int = 0
    buckets: list[dict] = field(default_factory=list)
    categories: list[dict] = field(default_factory=list)
    top_products: list[dict] = field(default_factory=list)
    staff: list[dict] = field(default_factory=list)


def _period(range_key: str, today: date) -> tuple[date, list[tuple[str, date, date]]]:
    """The first day and the chart buckets (label, first day, last day)."""
    if range_key == "7d":
        start = today - timedelta(days=6)
        return start, [((start + timedelta(days=n)).strftime("%a"), start + timedelta(days=n), start + timedelta(days=n)) for n in range(7)]
    if range_key == "30d":
        start = today - timedelta(days=29)
        buckets, first = [], start
        while first <= today:
            last = min(today, first + timedelta(days=6))
            buckets.append((first.strftime("%d %b"), first, last))
            first = last + timedelta(days=1)
        return start, buckets
    start = date(today.year, 1, 1)
    buckets = []
    for month in range(1, today.month + 1):
        first = date(today.year, month, 1)
        nxt = date(today.year + (month == 12), month % 12 + 1, 1)
        buckets.append((first.strftime("%b"), first, min(today, nxt - timedelta(days=1))))
    return start, buckets


def _day_col(column):
    # A literal (not a bound parameter), so SELECT and GROUP BY are the same expression to Postgres.
    return func.date(func.timezone(literal_column("'Africa/Kigali'"), column))


def build(db: Session, range_key: str) -> Report:
    today = shop_settings.today_in_kigali()
    start, buckets = _period(range_key, today)
    lo, hi = kigali_day_start(start), kigali_day_start(today + timedelta(days=1))
    report = Report(range=range_key, label=RANGES[range_key], start=start, end=today)
    live = (Order.status != "cancelled", Order.created_at >= lo, Order.created_at < hi)

    order_days = dict(db.execute(select(_day_col(Order.created_at), func.sum(Order.total_rwf)).where(*live)
                                 .group_by(_day_col(Order.created_at))).all())
    print_done = (PrintRequest.status == "collected", PrintRequest.paid_at >= lo, PrintRequest.paid_at < hi)
    print_days = dict(db.execute(select(_day_col(PrintRequest.paid_at), func.sum(PrintRequest.quote_rwf)).where(*print_done)
                                 .group_by(_day_col(PrintRequest.paid_at))).all())

    for label, first, last in buckets:
        total = sum(int(v or 0) for d, v in order_days.items() if first <= d <= last)
        total += sum(int(v or 0) for d, v in print_days.items() if first <= d <= last)
        report.buckets.append({"label": label, "start": first, "end": last, "revenue_rwf": total})

    count, order_total = db.execute(select(func.count(), func.coalesce(func.sum(Order.total_rwf), 0)).where(*live)).one()
    jobs, print_total = db.execute(select(func.count(), func.coalesce(func.sum(PrintRequest.quote_rwf), 0)).where(*print_done)).one()
    report.orders = count
    report.print_jobs_completed = jobs
    report.revenue_rwf = int(order_total) + int(print_total)
    report.average_order_rwf = int(order_total) // count if count else 0

    line_total = func.sum(OrderItem.unit_price_rwf * OrderItem.qty)
    by_category = db.execute(
        select(Category.name, Category.tint, line_total).select_from(OrderItem)
        .join(Order, Order.id == OrderItem.order_id).join(Product, Product.id == OrderItem.product_id)
        .join(Category, Category.id == Product.category_id).where(*live)
        .group_by(Category.name, Category.tint).order_by(line_total.desc())
    ).all()
    items_total = sum(int(r[2]) for r in by_category) or 1
    report.categories = [{"name": n, "tint": t, "revenue_rwf": int(v), "share": round(int(v) * 100 / items_total)} for n, t, v in by_category]
    if print_total:
        report.categories.append({"name": "Print & binding", "tint": "blush", "revenue_rwf": int(print_total), "share": None})

    top = db.execute(
        select(OrderItem.name_snapshot, func.sum(OrderItem.qty), line_total).select_from(OrderItem)
        .join(Order, Order.id == OrderItem.order_id).where(*live)
        .group_by(OrderItem.name_snapshot).order_by(line_total.desc()).limit(5)
    ).all()
    report.top_products = [{"name": n, "units": int(u), "revenue_rwf": int(v)} for n, u, v in top]
    report.staff = _staff_performance(db, lo, hi)
    return report


def _staff_performance(db: Session, lo: datetime, hi: datetime) -> list[dict]:
    """Orders handled = orders whose status they changed + walk-ins they recorded. Print jobs = requests they
    moved on. Sales = their walk-ins + online orders they marked collected."""
    handled: dict = {}
    for staff_id, order_id in db.execute(
        select(OrderStatusHistory.staff_id, OrderStatusHistory.order_id)
        .where(OrderStatusHistory.staff_id.is_not(None), OrderStatusHistory.event.is_(None),
               OrderStatusHistory.at >= lo, OrderStatusHistory.at < hi)
    ):
        handled.setdefault(staff_id, set()).add(order_id)
    sales: dict = {}
    for staff_id, total in db.execute(
        select(Order.created_by_staff_id, func.sum(Order.total_rwf))
        .where(Order.source == "walk_in", Order.status != "cancelled", Order.created_at >= lo, Order.created_at < hi)
        .group_by(Order.created_by_staff_id)
    ):
        sales[staff_id] = sales.get(staff_id, 0) + int(total)
    for staff_id, total in db.execute(
        select(OrderStatusHistory.staff_id, func.sum(Order.total_rwf)).join(Order, Order.id == OrderStatusHistory.order_id)
        .where(OrderStatusHistory.to_status == "collected", OrderStatusHistory.event.is_(None), Order.source == "online",
               OrderStatusHistory.at >= lo, OrderStatusHistory.at < hi)
        .group_by(OrderStatusHistory.staff_id)
    ):
        sales[staff_id] = sales.get(staff_id, 0) + int(total)
    jobs: dict = {}
    for staff_id, request_id in db.execute(
        select(PrintRequestHistory.staff_id, PrintRequestHistory.request_id)
        .where(PrintRequestHistory.staff_id.is_not(None), PrintRequestHistory.event.is_(None),
               PrintRequestHistory.at >= lo, PrintRequestHistory.at < hi)
    ):
        jobs.setdefault(staff_id, set()).add(request_id)

    people = db.scalars(select(Staff).order_by(Staff.name)).all()
    rows = []
    for person in people:
        if not person.is_active and person.id not in handled and person.id not in sales:
            continue
        rows.append({"name": person.name, "role": person.role, "orders_handled": len(handled.get(person.id, ())),
                     "print_jobs": len(jobs.get(person.id, ())), "sales_rwf": sales.get(person.id, 0)})
    return sorted(rows, key=lambda r: (-r["sales_rwf"], r["name"]))


# Exports


def _rows(report: Report) -> dict[str, list[list]]:
    return {
        "Summary": [["Range", f"{report.label} ({report.start} to {report.end})"], ["Revenue (RWF)", report.revenue_rwf],
                    ["Orders", report.orders], ["Average order (RWF)", report.average_order_rwf],
                    ["Print jobs completed", report.print_jobs_completed]],
        "Revenue": [["Period", "From", "To", "Revenue (RWF)"]] + [[b["label"], str(b["start"]), str(b["end"]), b["revenue_rwf"]] for b in report.buckets],
        "Categories": [["Category", "Revenue (RWF)", "Share %"]] + [[c["name"], c["revenue_rwf"], c["share"] if c["share"] is not None else ""] for c in report.categories],
        "Top products": [["Product", "Units", "Revenue (RWF)"]] + [[t["name"], t["units"], t["revenue_rwf"]] for t in report.top_products],
        "Staff": [["Staff", "Role", "Orders handled", "Print jobs", "Sales (RWF)"]] + [[s["name"], s["role"], s["orders_handled"], s["print_jobs"], s["sales_rwf"]] for s in report.staff],
    }


def to_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    book.remove(book.active)
    for title, rows in _rows(report).items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in column) + 2))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _latin1(text: object) -> str:
    """The PDF's built-in font only covers Latin-1; swap the few characters the app uses outside it."""
    return str(text).replace("–", "-").replace("—", "-").replace("×", "x").replace("→", "->").encode("latin-1", "replace").decode("latin-1")


def to_pdf(report: Report, shop_name: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, _latin1(f"{shop_name} - report, {report.label}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, _latin1(f"{report.start} to {report.end} (Kigali time)"), new_x="LMARGIN", new_y="NEXT")
    for title, rows in _rows(report).items():
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 13)
        pdf.cell(0, 8, _latin1(title), new_x="LMARGIN", new_y="NEXT")
        width = (pdf.w - pdf.l_margin - pdf.r_margin) / max(len(r) for r in rows)
        for i, row in enumerate(rows):
            pdf.set_font("Helvetica", "B" if i == 0 and title != "Summary" else "", 10)
            for value in row:
                text = f"{value:,}" if isinstance(value, int) else _latin1(value)
                pdf.cell(width, 7, text[:48], border="B")
            pdf.ln()
    return bytes(pdf.output())
